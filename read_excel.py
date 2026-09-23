import openpyxl
import os

base = r"C:\Users\Utente\Desktop\documenti elia (appunti)\fantacalcio\fanta27\coach-beard-starter 2\coach-beard-starter"

files = [
    os.path.join(base, 'Dati/Dati Su Giocatori, Allenatori e Squadre/Guida_Asta_202627.xlsx'),
    os.path.join(base, 'Dati/Listone E quotazioni/Quotazioni_Fantacalcio_Stagione_2026_27.xlsx'),
    os.path.join(base, 'Dati/Strategie/StrategiaFanta.xlsx'),
    os.path.join(base, 'Dati/Strategie/Pupilli.xlsx'),
]

for f in files:
    print('='*90)
    print('FILE:', os.path.relpath(f, base))
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        print('SHEETS:', wb.sheetnames)
        for sn in wb.sheetnames:
            ws = wb[sn]
            print(f'  Sheet [{sn}]: {ws.max_row} rows x {ws.max_column} cols')
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 6: break
                print('   ROW', i, ':', row)
            print()
    except Exception as e:
        print('  ERROR:', type(e).__name__, e)
